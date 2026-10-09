"""Extracts readable text from downloaded attachment bytes. Supports the
formats that actually show up in real ingested mail in meaningful volume
(checked live against real data before choosing this list): pdf, docx,
xlsx, csv, and (2026-09-10) png/jpeg/gif/webp images via Claude vision --
a scanned document, screenshot, or photo of a page is otherwise invisible
to every downstream tool exactly the way a truncated attachment was
before that same day's fix. Legacy .doc/.xls (pre-2007 binary Office
formats) remain explicitly out of scope -- would need a different library
for a handful of real instances, not worth it. Unsupported formats
degrade to a clear, honest "not supported" note rather than silently
doing nothing.

Never raises -- a single bad attachment (corrupted file, unexpected
internal structure, password-protected) degrades to a short bracketed
note stored in the same place real extracted text would go, so a caller
never needs a separate try/except and a query for "did extraction
actually happen" is just "is extracted_text NULL" (not attempted, e.g.
an unsupported binary format skipped before download) vs "here's what
happened" (a real value, success or an explained failure).
"""
import csv
import io
import logging

logger = logging.getLogger("ingestion.attachment_extract")

# 2026-09-10: removed the storage-time character cap that used to live here
# (was 50000, raised once to 500000, then removed outright). It silently
# dropped an xlsx's later sheets -- sheet order in the file, not relevance,
# determines what survives a truncated slice -- costing a real payroll
# query its August 2026 tab even though earlier sheets extracted fine.
# Extraction now always stores the complete text; bounding what a single
# get_attachment_text call returns is mcp_server/server.py's job, via
# sheet-selection and pagination, not something decided at ingestion time.

SUPPORTED_EXTENSIONS = {"pdf", "docx", "xlsx", "csv", "png", "jpg", "jpeg", "gif", "webp"}

MIN_IMAGE_BYTES = 10000  # below this, almost certainly a logo, icon, or tracking pixel --
# not a genuine document worth spending a vision API call transcribing. Real scanned
# documents/screenshots/photos ingested in practice all landed well above this.


def _extract_pdf(data):
    import pypdf
    reader = pypdf.PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"--- page {i + 1} ---\n{text}")
    return "\n\n".join(pages)


def _extract_docx(data):
    import docx
    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_xlsx(data):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    sections = []
    for sheet in wb.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            if any(c is not None for c in row):
                rows.append(" | ".join("" if c is None else str(c) for c in row))
        if rows:
            sections.append(f"--- sheet: {sheet.title} ---\n" + "\n".join(rows))
    return "\n\n".join(sections)


def _extract_csv(data):
    text = data.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    return "\n".join(" | ".join(row) for row in reader)


def _detect_image_media_type(data):
    """Sniffs the real format from magic bytes rather than trusting the
    filename extension (which can be wrong -- a mislabeled .jpg is a real
    thing) -- cheap and exact for the four formats Claude's vision API
    accepts. None means "not actually one of these", handled by the caller
    as an honest unsupported-format note rather than a failed API call."""
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return "image/png"
    if data[:3] == b"\xff\xd8\xff":
        return "image/jpeg"
    if data[:6] in (b"GIF87a", b"GIF89a"):
        return "image/gif"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp"
    return None


def _extract_image(data):
    """Claude vision, not a traditional OCR library -- this codebase already
    runs every other LLM-judgment task (tier classification, redaction,
    wiki-drafting) through the Anthropic API, so this adds zero new vendor
    surface, and vision handles messy real-world photos/screenshots better
    than character-recognition-only OCR. Sonnet, not Haiku -- misreading a
    document (a wrong number, a garbled name) is worse than not extracting
    it at all, unlike a coarse classification call where being roughly
    right is fine. Synchronous client, matching every other per-item caller
    in this codebase (merge_live_duplicates.py, inline_tier_classifier.py)
    -- extract_text() is called synchronously per attachment, not batched,
    so there is no event loop here to make async worth the complexity."""
    import base64
    import os

    import anthropic

    media_type = _detect_image_media_type(data)
    if media_type is None:
        return "[unrecognized image format -- not one of png/jpeg/gif/webp]"
    if len(data) < MIN_IMAGE_BYTES:
        return "[image too small to be a meaningful document -- likely a logo, icon, or tracking pixel, skipped]"

    client = anthropic.Anthropic(api_key=os.environ["CLASSIFIER_ANTHROPIC_API_KEY"])
    resp = client.messages.create(
        model="claude-sonnet-5",
        max_tokens=4096,
        messages=[{
            "role": "user",
            "content": [
                {"type": "image", "source": {"type": "base64", "media_type": media_type,
                                              "data": base64.standard_b64encode(data).decode()}},
                {"type": "text", "text": (
                    "Transcribe all readable text from this image exactly as it appears, "
                    "preserving table structure (separate columns with ' | ') if present. "
                    "If this image is a logo, icon, decorative graphic, or otherwise contains "
                    "no meaningful readable text, respond with exactly: [no meaningful text content]"
                )},
            ],
        }],
    )
    text = "".join(block.text for block in resp.content if block.type == "text").strip()
    return text or "[no meaningful text content]"


_EXTRACTORS = {
    "pdf": _extract_pdf,
    "docx": _extract_docx,
    "xlsx": _extract_xlsx,
    "csv": _extract_csv,
    "png": _extract_image,
    "jpg": _extract_image,
    "jpeg": _extract_image,
    "gif": _extract_image,
    "webp": _extract_image,
}


def extractable(filename):
    """Cheap pre-check, called before spending an API call downloading
    the actual bytes -- no point fetching a binary this module can't do
    anything with."""
    ext = (filename or "").rsplit(".", 1)[-1].lower() if "." in (filename or "") else ""
    return ext in SUPPORTED_EXTENSIONS


def extract_text(filename, data):
    """Returns extracted text, or a short bracketed explanatory note --
    never raises, never returns None or empty."""
    ext = (filename or "").rsplit(".", 1)[-1].lower() if "." in (filename or "") else ""
    extractor = _EXTRACTORS.get(ext)
    if not extractor:
        if ext in ("xls", "doc"):
            return f"[legacy .{ext} format not supported for extraction -- re-save as .{ext}x to enable this]"
        return f"[unsupported format for text extraction: .{ext or 'unknown'}]"
    try:
        text = extractor(data)
        # Postgres text columns reject embedded NUL bytes outright ("A
        # string literal cannot contain NUL (0x00) characters") -- some
        # malformed/corrupted source PDFs decode to text containing them.
        # Stripping here (not just at the DB-write call site) means every
        # caller gets a value that's always safely storable.
        text = text.replace("\x00", "")
        if not text.strip():
            return "[no extractable text found -- file may be empty, scanned/image-based, or password-protected]"
        return text
    except Exception as e:
        logger.warning("attachment extraction failed for %s: %s", filename, e)
        return f"[extraction failed: {type(e).__name__}]"
